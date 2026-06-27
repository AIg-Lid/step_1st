# -*- coding: utf-8 -*-
"""
导出路由模块 - DOCX / XLSX 文件生成与下载
"""

from __future__ import annotations

import io
import urllib.parse
import logging
import re
from datetime import date

from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from docx import Document as DocxDocument
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

from app.models import ExportDocxRequest, ExportXlsxRequest, ErrorResponse

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api", tags=["Export API"])

# ═══════════════════════════════════════════════════════════════════
# Jeffrey 品牌配色 & 常量
# ═══════════════════════════════════════════════════════════════════

NAVY = RGBColor(0x10, 0x2A, 0x54)
CYAN = RGBColor(0x00, 0xB4, 0xD8)
ORANGE = RGBColor(0xFF, 0x7A, 0x1A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x55, 0x60, 0x7A)
MIDGRAY = RGBColor(0xD8, 0xE0, 0xEA)
LIGHT_BG = "EAF1FA"
NAVY_HEX = "102A54"
CYAN_HEX = "00B4D8"
ORANGE_HEX = "FF7A1A"


# ═══════════════════════════════════════════════════════════════════
# POST /api/export_docx
# ═══════════════════════════════════════════════════════════════════

@router.post(
    "/export_docx",
    summary="导出 Word 文档 (.docx)",
    description="接收文章标题和内容块列表，生成 Jeffrey 品牌风格的 .docx 文件",
    responses={
        200: {"description": "DOCX 文件", "content": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document": {}}},
        400: {"model": ErrorResponse},
    },
)
async def export_docx(body: ExportDocxRequest):
    """导出 Word 文档"""
    if not body.content:
        raise HTTPException(status_code=400, detail={"error": "bad_request", "detail": "content 不能为空"})

    doc = DocxDocument()

    # ── 样式配置 ──
    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.5

    for level, (size, color) in enumerate([(24, NAVY), (18, CYAN), (14, CYAN)], 1):
        hname = f'Heading {level}'
        hs = doc.styles[hname]
        hs.font.name = 'Arial'
        hs.font.size = Pt(size)
        hs.font.bold = True
        hs.font.color.rgb = color
        hs.paragraph_format.space_before = Pt(18 if level == 1 else 12)
        hs.paragraph_format.space_after = Pt(8)

    # ── 页边距 ──
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)

    # ── 封面 ──
    for _ in range(6):
        doc.add_paragraph('')

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(body.title)
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = NAVY
    run.font.name = 'Arial'

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('Jeffrey_AI_step_1st 智能办公平台')
    run.font.size = Pt(12)
    run.font.color.rgb = GRAY
    run.font.italic = True
    run.font.name = 'Arial'

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f'生成日期：{date.today().strftime("%Y-%m-%d")}')
    run.font.size = Pt(10)
    run.font.color.rgb = GRAY
    run.font.name = 'Arial'

    doc.add_page_break()

    # ── 正文 ──
    for block in body.content:
        t = block.type
        text = block.text

        if t in ('heading1', 'heading2', 'heading3'):
            h = doc.add_heading(text, level=int(t[-1]))
        elif t == 'bullet':
            p = doc.add_paragraph(text, style='List Bullet')
        else:
            p = doc.add_paragraph(text)

    # ── 页脚 ──
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f'— 本文档由 {body.author} 自动生成 —')
    run.font.size = Pt(9)
    run.font.color.rgb = GRAY
    run.font.italic = True

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    fname = re.sub(r'[\\/:*?"<>|]', '_', body.title)[:60]
    logger.info("DOCX导出: title='%s', blocks=%d", body.title[:30], len(body.content))

    encoded = urllib.parse.quote(fname + '.docx')
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded}"},
    )


# ═══════════════════════════════════════════════════════════════════
# POST /api/export_xlsx
# ═══════════════════════════════════════════════════════════════════

@router.post(
    "/export_xlsx",
    summary="导出 Excel 表格 (.xlsx)",
    description="接收表头和行数据，生成 Jeffrey 品牌风格的 .xlsx 文件",
    responses={
        200: {"description": "XLSX 文件", "content": {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": {}}},
        400: {"model": ErrorResponse},
    },
)
async def export_xlsx(body: ExportXlsxRequest):
    """导出 Excel 表格"""
    if not body.headers:
        raise HTTPException(status_code=400, detail={"error": "bad_request", "detail": "headers 不能为空"})

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    # ── 标题行 ──
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(body.headers))
    title_cell = ws.cell(row=1, column=1, value=body.title)
    title_cell.font = Font(name='Arial', bold=True, color=NAVY_HEX, size=14)
    title_cell.alignment = Alignment(horizontal='left', vertical='center')

    # ── 副标题 ──
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(body.headers))
    sub_cell = ws.cell(row=2, column=1, value=f'Jeffrey_AI_step_1st · 生成日期：{date.today().strftime("%Y-%m-%d")}')
    sub_cell.font = Font(name='Arial', italic=True, color=GRAY.hex if hasattr(GRAY, 'hex') else '55607A', size=9)
    sub_cell.alignment = Alignment(horizontal='left', vertical='center')

    # ── 表头 ──
    header_font = Font(name='Arial', bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill('solid', fgColor=NAVY_HEX)
    thin_border = Border(
        left=Side(style='thin', color=MIDGRAY.hex if hasattr(MIDGRAY, 'hex') else 'D8E0EA'),
        right=Side(style='thin', color=MIDGRAY.hex if hasattr(MIDGRAY, 'hex') else 'D8E0EA'),
        top=Side(style='thin', color=MIDGRAY.hex if hasattr(MIDGRAY, 'hex') else 'D8E0EA'),
        bottom=Side(style='thin', color=MIDGRAY.hex if hasattr(MIDGRAY, 'hex') else 'D8E0EA'),
    )
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)
    data_font = Font(name='Arial', color='1F293A', size=10)
    even_fill = PatternFill('solid', fgColor=LIGHT_BG)

    for col_idx, header in enumerate(body.headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = max(16, len(header) * 2.5 + 4)

    # ── 数据行 ──
    for row_idx, row_data in enumerate(body.rows, 5):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = data_font
            cell.alignment = left_align if col_idx == 1 else center_align
            cell.border = thin_border
            if row_idx % 2 == 0:
                cell.fill = even_fill

    ws.freeze_panes = 'A5'

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    fname = re.sub(r'[\\/:*?"<>|]', '_', body.title)[:60]
    logger.info("XLSX导出: title='%s', rows=%d, cols=%d", body.title[:30], len(body.rows), len(body.headers))

    encoded = urllib.parse.quote(fname + '.xlsx')
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded}"},
    )
